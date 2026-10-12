#!/usr/bin/env python3
"""依 JSON 大綱，用公司模板 (assets/template.pptx) 產生新簡報。

用法: python3 build_deck.py spec.json out.pptx [--template path.pptx]

spec.json 格式:
{
  "slides": [
    {"layout": "cover",
     "title": "主標題第一行\\n第二行",          # idx 0
     "subtitle": "副標題",                     # idx 10
     "presenter": "報告人：部門_姓名",          # idx 16
     "date": "2026.10.12"},                    # idx 18
    {"layout": "content",                      # = 'Content 2'
     "title": "頁面標題",
     "lead": "導言（前面會自動有橘紅色方塊）",
     "body": ["## 段落標題", "**關鍵詞**：說明", "第二點", {"text": "子項", "level": 1}],
     "notes": "講稿"},
    {"layout": "Content 6", "placeholders": {"16": ["..."], "17": {"image": "fig.png"}}},
    {"layout": "thanks"}
  ]
}

- layout 可用別名 (cover / content / two-content / dashboard / thanks) 或模板中的版型名稱。
- title / lead / body / footer 分別對應 idx 0 / 15 / 2 / 3；其他欄位用 "placeholders": {idx: 值}。
- 值: 字串（\\n 分段）、清單（每項一段）、或 {"image": 路徑}。
- 清單中以 "## " 開頭的項目 = 粗體段落標題（不編號）；"**文字**" = 粗體片段。
- 未填寫的 placeholder 會自動刪除，避免出現「按一下以新增文字」。
"""
import argparse
import json
import re
import sys
from pathlib import Path

from pptx import Presentation

HERE = Path(__file__).resolve().parent.parent
DEFAULT_TEMPLATE = HERE / "assets" / "template.pptx"

ALIASES = {
    "cover": "(一般內外部使用)封面",
    "cover-blue": "（對外使用）封面_高階主管/公司介紹使用_藍色",
    "toc": "目錄",
    "section": "（通用）過場頁",
    "content": "Content 2",
    "two-content": "Content 6",
    "dashboard": "Content 43",
    "table": "內頁（表格呈現）",
    "thanks": "（通用）感謝頁",
}
FIELD_IDX = {"title": 0, "lead": 15, "body": 2, "footer": 3,
             "subtitle": 10, "presenter": 16, "date": 18}
BOLD_RE = re.compile(r"(\*\*.+?\*\*)")


def find_layout(prs, name):
    name = ALIASES.get(name, name)
    for layout in prs.slide_layouts:
        if layout.name == name:
            return layout
    names = ", ".join(l.name for l in prs.slide_layouts)
    sys.exit(f"找不到版型 {name!r}。可用: {names}")


def add_runs(paragraph, text):
    for part in BOLD_RE.split(text):
        if not part:
            continue
        run = paragraph.add_run()
        if part.startswith("**") and part.endswith("**"):
            run.text = part[2:-2]
            run.font.bold = True
        else:
            run.text = part


def no_bullet(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    for tag in ("a:buNone", "a:buAutoNum", "a:buChar"):
        for el in pPr.findall("{http://schemas.openxmlformats.org/drawingml/2006/main}" + tag.split(":")[1]):
            pPr.remove(el)
    pPr.insert(0, pPr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}buNone", {}))


def fill_text(shape, value):
    items = value.split("\n") if isinstance(value, str) else list(value)
    tf = shape.text_frame
    tf.clear()
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if isinstance(item, dict):
            p.level = item.get("level", 0)
            text = item["text"]
            if item.get("bold"):
                text = f"**{text}**"
        else:
            text = item
        if text.startswith("## "):
            add_runs(p, f"**{text[3:]}**")
            no_bullet(p)
        else:
            add_runs(p, text)


def fill_image(slide, shape, path):
    left, top, width, height = shape.left, shape.top, shape.width, shape.height
    if hasattr(shape, "insert_picture"):
        shape.insert_picture(str(path))
        return
    pic = slide.shapes.add_picture(str(path), left, top)
    # 等比縮放塞進 placeholder 範圍並置中
    scale = min(width / pic.width, height / pic.height)
    pic.width, pic.height = int(pic.width * scale), int(pic.height * scale)
    pic.left = left + (width - pic.width) // 2
    pic.top = top + (height - pic.height) // 2
    shape._element.getparent().remove(shape._element)


def build(spec, out, template, base_dir):
    prs = Presentation(template)
    for s in spec["slides"]:
        layout = find_layout(prs, s["layout"])
        slide = prs.slides.add_slide(layout)
        values = {str(idx): s[k] for k, idx in FIELD_IDX.items() if k in s}
        values.update({str(k): v for k, v in s.get("placeholders", {}).items()})
        by_idx = {str(ph.placeholder_format.idx): ph for ph in slide.placeholders}
        for idx, value in values.items():
            ph = by_idx.get(idx)
            if ph is None:
                print(f"警告: 版型 {layout.name!r} 沒有 idx={idx}，已略過", file=sys.stderr)
                continue
            if isinstance(value, dict) and "image" in value:
                fill_image(slide, ph, (base_dir / value["image"]).resolve())
            else:
                fill_text(ph, value)
        for idx, ph in by_idx.items():
            if idx not in values:
                ph._element.getparent().remove(ph._element)
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = s["notes"]
    prs.save(out)
    print(f"已輸出 {out}（{len(spec['slides'])} 頁）")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("spec")
    ap.add_argument("out")
    ap.add_argument("--template", default=DEFAULT_TEMPLATE)
    a = ap.parse_args()
    spec_path = Path(a.spec)
    build(json.loads(spec_path.read_text(encoding="utf-8")), a.out, a.template, spec_path.parent)


if __name__ == "__main__":
    main()
